"""Laporan: rekap transaksi & durasi parkir — Excel, CSV & PDF (PRD §5 /admin/laporan)."""
import csv
import io
from datetime import date, timedelta
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.deps import DB, AdminUser
from app.core.timeutil import day_bounds, today, tz
from app.core.config import get_settings
from app.models import Incident, Transaction
from app.services import scope, storage
from app.services import transactions as tx_svc

router = APIRouter(prefix="/admin/reports", tags=["reports"])

STATUS_LABEL = {
    "parked": "Masih parkir", "incident_pending": "Menunggu supervisor",
    "incident_confirmed": "Kerusakan dicatat", "released": "Selesai",
}
INCIDENT_LABEL = {"pending": "Menunggu", "confirmed": "Kerusakan baru", "dismissed": "Tidak ada kerusakan"}


class Period:
    def __init__(self, scope_: str = Query("all", alias="scope"),
                 date_from: date | None = Query(None, alias="from"), date_to: date | None = Query(None, alias="to")):
        self.scope = scope_
        self.d_from, self.d_to, self.start, self.end = _range(date_from, date_to)


PeriodQ = Annotated[Period, Depends()]


def _range(date_from: date | None, date_to: date | None):
    d_to = date_to or today()
    d_from = date_from or d_to - timedelta(days=6)
    if d_from > d_to:
        raise HTTPException(422, "Tanggal awal setelah tanggal akhir")
    if (d_to - d_from).days > 366:
        raise HTTPException(422, "Rentang maksimal 1 tahun")
    return d_from, d_to, day_bounds(d_from)[0], day_bounds(d_to)[1]


def _txs(db, user, scope_, start, end):
    ids = scope.resolve_scope(db, user, scope_)
    stmt = (select(Transaction)
            .options(selectinload(Transaction.post), selectinload(Transaction.site),
                     selectinload(Transaction.checkin_user), selectinload(Transaction.checkout_user),
                     selectinload(Transaction.incidents), selectinload(Transaction.checks),
                     selectinload(Transaction.inside_items))
            .where(Transaction.checkin_at >= start, Transaction.checkin_at < end)
            .order_by(Transaction.checkin_at))
    if ids is not None:
        stmt = stmt.where(Transaction.post_id.in_(ids or {-1}))
    return db.scalars(stmt).all()


def _fmt(dt) -> str:
    return dt.astimezone(tz()).strftime("%Y-%m-%d %H:%M") if dt else ""


def _minutes(tx) -> int | None:
    end = tx.released_at or tx.checkout_at
    return int((end - tx.checkin_at).total_seconds() // 60) if end else None


def _csv(name: str, header: list[str], rows) -> StreamingResponse:
    buf = io.StringIO()
    buf.write("﻿")  # supaya Excel membaca UTF-8
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return StreamingResponse(iter([buf.getvalue()]), media_type="text/csv; charset=utf-8",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


TX_HEADER = ["Nomor polisi", "Lokasi", "Pos", "Masuk", "Petugas masuk", "Keluar", "Petugas keluar",
             "Durasi (menit)", "Status", "Kondisi (selain mulus)", "Barang di dalam", "Catatan"]


def _local(dt):
    """Waktu lokal tanpa zona, supaya Excel menampilkannya sebagai tanggal-jam biasa."""
    return dt.astimezone(tz()).replace(tzinfo=None) if dt else None


def _tx_row(tx) -> list:
    return [tx.plate, tx.site.name, tx.post.name, _local(tx.checkin_at), tx.checkin_user.name,
            _local(tx.checkout_at), tx.checkout_user.name if tx.checkout_user else "", _minutes(tx),
            STATUS_LABEL[tx.status], _conditions(tx), _inside(tx), (tx.notes or "").replace("\n", " ")]


@router.get("/transactions.csv")
def transactions_csv(user: AdminUser, db: DB, p: PeriodQ):
    rows = []
    for tx in _txs(db, user, p.scope, p.start, p.end):
        r = _tx_row(tx)
        r[3], r[5] = _fmt(tx.checkin_at), _fmt(tx.checkout_at)
        r[7] = "" if r[7] is None else r[7]
        rows.append(r)
    return _csv(f"ceklis-transaksi-{p.d_from}-{p.d_to}.csv", TX_HEADER, rows)


def _sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(header)
    for r in rows:
        ws.append(r)
    head_fill = PatternFill("solid", fgColor="F5C518")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if hasattr(c.value, "hour"):
                c.number_format = "dd/mm/yyyy hh:mm"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


@router.get("/transactions.xlsx")
def transactions_xlsx(user: AdminUser, db: DB, p: PeriodQ):
    txs = _txs(db, user, p.scope, p.start, p.end)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi"
    _sheet(ws, TX_HEADER, [_tx_row(t) for t in txs], [14, 22, 16, 17, 18, 17, 18, 10, 13, 40, 28, 30])

    rekap = wb.create_sheet("Rekap per pos")
    _sheet(rekap, ["Lokasi", "Pos", "Check in", "Selesai", "Masih parkir", "Rata-rata durasi (menit)"],
           [[site, post, n, done, n - done, avg] for site, post, n, done, avg in _per_post(txs)], [24, 18, 10, 10, 13, 22])

    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="ceklis-transaksi-{p.d_from}-{p.d_to}.xlsx"'})


def _per_post(txs) -> list[tuple[str, str, int, int, int | None]]:
    by_loc: dict[tuple[str, str], list] = {}
    for t in txs:
        by_loc.setdefault((t.site.name, t.post.name), []).append(t)
    out = []
    for (site, post), items in sorted(by_loc.items()):
        ds = [m for m in (_minutes(t) for t in items) if m is not None]
        out.append((site, post, len(items), sum(t.status == "released" for t in items), sum(ds) // len(ds) if ds else None))
    return out


def _conditions(tx) -> str:
    labels = {c["id"]: c["label"] for c in (tx.template_snapshot or {}).get("conditions", [])}
    bad = [c for c in tx.checks if c.phase == "in" and not c.checked]
    return ", ".join(f"{c.label_snapshot}: {labels.get(c.condition or '', 'perlu dicek')}" for c in bad) or "Semua baik"


def _inside(tx) -> str:
    items = [i for i in tx.inside_items if i.phase == "in" and i.present]
    return ", ".join(f"{i.label} ({i.extra_text})" if i.extra_text else i.label for i in items) or "Tidak ada"


@router.get("/incidents.csv")
def incidents_csv(user: AdminUser, db: DB, p: PeriodQ):
    ids = scope.resolve_scope(db, user, p.scope)
    stmt = (select(Incident).join(Transaction)
            .options(selectinload(Incident.transaction).selectinload(Transaction.post),
                     selectinload(Incident.transaction).selectinload(Transaction.site),
                     selectinload(Incident.reporter), selectinload(Incident.decider))
            .where(Incident.created_at >= p.start, Incident.created_at < p.end).order_by(Incident.created_at))
    if ids is not None:
        stmt = stmt.where(Transaction.post_id.in_(ids or {-1}))
    rows = [[
        i.code, i.transaction.ticket_code, i.transaction.plate, i.transaction.site.name, i.transaction.post.name,
        _fmt(i.created_at), i.reporter.name, ", ".join(i.parts), (i.note or "").replace("\n", " "),
        INCIDENT_LABEL[i.status], i.decider.name if i.decider else "", _fmt(i.decided_at),
        (i.decision_note or "").replace("\n", " "),
    ] for i in db.scalars(stmt)]
    return _csv(f"ceklis-insiden-{p.d_from}-{p.d_to}.csv", [
        "Kode", "Tiket cek list", "Nomor polisi", "Lokasi", "Pos", "Dilaporkan", "Petugas", "Bagian",
        "Catatan petugas", "Keputusan", "Supervisor", "Diputuskan", "Catatan pemeriksaan",
    ], rows)


def _dur(m: int | None) -> str:
    if m is None:
        return "-"
    return f"{m // 60}j {m % 60}m" if m >= 60 else f"{m}m"


@router.get("/summary.pdf")
def summary_pdf(user: AdminUser, db: DB, p: PeriodQ):
    """Laporan PDF: ringkasan + rekap per pos + rincian setiap transaksi."""
    d_from, d_to = p.d_from, p.d_to
    txs = _txs(db, user, p.scope, p.start, p.end)
    durations = [m for m in (_minutes(t) for t in txs) if m is not None]

    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=7.5, leading=9)
    buf = io.BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title="Laporan Ceklis")
    ink, line, accent = colors.HexColor("#0E100F"), colors.HexColor("#DCDDD6"), colors.HexColor("#F5C518")

    def table(data, widths):
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), ink), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F3")]),
        ]))
        return t

    def para(*lines) -> Paragraph:
        esc = (str(x).replace("&", "&amp;").replace("<", "&lt;") for x in lines)
        return Paragraph("<br/>".join(esc), cell)

    story = [
        Paragraph("<b>ceklis</b> — Laporan transaksi", styles["Title"]),
        Paragraph(f"Periode {d_from:%d %b %Y} – {d_to:%d %b %Y}", styles["Normal"]),
        Spacer(1, 5 * mm),
    ]
    kpi = Table([
        ["Check in", "Selesai", "Masih parkir", "Rata-rata durasi"],
        [len(txs), sum(t.status == "released" for t in txs), sum(t.status != "released" for t in txs),
         _dur(sum(durations) // len(durations) if durations else None)],
    ], colWidths=[45 * mm] * 4)
    kpi.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), accent), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, 1), 14), ("GRID", (0, 0), (-1, -1), 0.4, line),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 1), (-1, 1), 6), ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
    ]))
    story += [kpi, Spacer(1, 5 * mm), Paragraph("<b>Per lokasi / pos</b>", styles["Heading3"])]
    loc_rows = [["Lokasi", "Pos", "Check in", "Selesai", "Masih parkir", "Rata-rata durasi"]]
    loc_rows += [[site, post, n, done, n - done, _dur(avg)] for site, post, n, done, avg in _per_post(txs)]
    story.append(table(loc_rows, [70 * mm, 45 * mm, 25 * mm, 25 * mm, 28 * mm, 35 * mm]))

    story += [Spacer(1, 6 * mm), Paragraph(f"<b>Rincian transaksi</b> ({len(txs)})", styles["Heading3"])]
    rows = [["No", "Nomor polisi", "Lokasi · Pos", "Masuk", "Keluar", "Durasi", "Status",
             "Kondisi (selain mulus)", "Barang di dalam"]]
    for n, t in enumerate(txs, start=1):
        rows.append([
            n, para(t.plate), para(f"{t.site.name} · {t.post.name}"),
            para(_fmt(t.checkin_at), t.checkin_user.name),
            para(_fmt(t.checkout_at), t.checkout_user.name) if t.checkout_user else para("-"),
            _dur(_minutes(t)), para(STATUS_LABEL[t.status]), para(_conditions(t)), para(_inside(t)),
        ])
    if len(rows) == 1:
        rows.append(["", "Tidak ada transaksi dalam periode ini.", "", "", "", "", "", "", ""])
    story.append(table(rows, [10 * mm, 25 * mm, 40 * mm, 30 * mm, 30 * mm, 16 * mm, 20 * mm, 58 * mm, 44 * mm]))

    doc.build(story)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="ceklis-laporan-{d_from}-{d_to}.pdf"'})


SIDE_NAME = {"depan": "Depan", "belakang": "Belakang", "kiri": "Kiri", "kanan": "Kanan"}


def _photo_image(key: str, width: float) -> Image | None:
    """Ambil foto dari storage, perkecil untuk PDF. None kalau tidak bisa dimuat."""
    from PIL import Image as PILImage

    try:
        raw = storage.internal_client().get_object(Bucket=get_settings().s3_bucket, Key=key)["Body"].read()
        im = PILImage.open(io.BytesIO(raw))
        im = im.convert("RGB")
        im.thumbnail((1000, 1000))
        out = io.BytesIO()
        im.save(out, "JPEG", quality=80)
        out.seek(0)
        return Image(out, width=width, height=width * im.height / im.width)
    except Exception:  # noqa: BLE001 — foto rusak/hilang tidak boleh menggagalkan PDF
        return None


@router.get("/transactions/{tx_id}.pdf")
def transaction_pdf(tx_id: int, user: AdminUser, db: DB):
    """PDF satu transaksi: data masuk/keluar, kondisi tiap bagian, barang di dalam, catatan, dan foto."""
    tx = tx_svc.get_transaction(db, str(tx_id))
    if not tx or not scope.can_access_post(db, user, tx.post_id):
        raise HTTPException(404, "Transaksi tidak ditemukan")

    styles = getSampleStyleSheet()
    cell = styles["BodyText"].clone("cell", fontSize=8.5, leading=11)
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9, textColor=colors.HexColor("#5B605A"))
    ink, line, accent = colors.HexColor("#0E100F"), colors.HexColor("#DCDDD6"), colors.HexColor("#F5C518")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=f"Ceklis {tx.plate}")
    full = A4[0] - 28 * mm

    def esc(x) -> str:
        return str(x).replace("&", "&amp;").replace("<", "&lt;")

    def grid(data, widths, header=True):
        t = Table(data, colWidths=widths, repeatRows=1 if header else 0)
        style = [("FONTSIZE", (0, 0), (-1, -1), 8.5), ("GRID", (0, 0), (-1, -1), 0.4, line),
                 ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if header:
            style += [("BACKGROUND", (0, 0), (-1, 0), ink), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                      ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        t.setStyle(TableStyle(style))
        return t

    labels = {c["id"]: c["label"] for c in (tx.template_snapshot or {}).get("conditions", [])}
    story = [
        Paragraph(f"<b>ceklis</b> — Laporan kendaraan <b>{esc(tx.plate)}</b>", styles["Title"]),
        Paragraph(f"{esc(tx.site.name)} · {esc(tx.post.name)} · {STATUS_LABEL[tx.status]}", styles["Normal"]),
        Spacer(1, 5 * mm),
    ]
    info = [
        ["Masuk", f"{_fmt(tx.checkin_at)} · oleh {tx.checkin_user.name}"],
        ["Keluar", f"{_fmt(tx.checkout_at)} · oleh {tx.checkout_user.name}" if tx.checkout_user else "Belum keluar"],
        ["Durasi", _dur(_minutes(tx))],
        ["Barang di dalam", _inside(tx)],
    ]
    if tx.notes:
        info.append(["Catatan masuk", tx.notes])
    if tx.checkout_notes:
        info.append(["Catatan keluar", tx.checkout_notes])
    info_t = Table([[Paragraph(f"<b>{k}</b>", cell), Paragraph(esc(v), cell)] for k, v in info],
                   colWidths=[40 * mm, full - 40 * mm])
    info_t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, line), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F7F7F3"))]))
    story += [info_t, Spacer(1, 5 * mm), Paragraph("<b>Kondisi bagian luar</b>", styles["Heading3"])]

    def cond(c) -> str:
        if c is None:
            return "-"
        return labels.get(c.condition or "", "Baik" if c.checked else "Perlu dicek")

    checks_in = [c for c in sorted(tx.checks, key=lambda c: c.position) if c.phase == "in"]
    out_by_id = {c.item_id: c for c in tx.checks if c.phase == "out"}
    rows = [["Bagian", "Grup", "Masuk", "Keluar"]]
    bad_rows = []
    for n, c in enumerate(checks_in, start=1):
        rows.append([c.label_snapshot, c.group_snapshot, cond(c), cond(out_by_id.get(c.item_id))])
        if not c.checked:
            bad_rows.append(n)
    t = grid(rows, [62 * mm, 42 * mm, (full - 104 * mm) / 2, (full - 104 * mm) / 2])
    t.setStyle(TableStyle([("TEXTCOLOR", (2, r), (2, r), colors.HexColor("#A3440B")) for r in bad_rows]
                          + [("FONTNAME", (2, r), (2, r), "Helvetica-Bold") for r in bad_rows]))
    story.append(t)

    photos = [ph for ph in tx.photos if ph.uploaded_at and not ph.deleted_at]
    if photos:
        names = {c.item_id: c.label_snapshot for c in tx.checks} | {i.item_id: i.label for i in tx.inside_items}
        story += [Spacer(1, 5 * mm), Paragraph(f"<b>Foto</b> ({len(photos)})", styles["Heading3"])]
        col_w = (full - 8 * mm) / 3
        cells = []
        for ph in sorted(photos, key=lambda x: (x.phase != "in", x.taken_at)):
            img = _photo_image(ph.object_key, col_w - 4 * mm)
            what = names.get(ph.item_id or "", SIDE_NAME.get(ph.side, ph.side))
            caption = Paragraph(f"{'Masuk' if ph.phase == 'in' else 'Keluar'} · {esc(what)}<br/>{_fmt(ph.taken_at)}", small)
            cells.append([img or Paragraph("Foto tidak bisa dimuat", small), caption])
        rows = [cells[i:i + 3] for i in range(0, len(cells), 3)]
        rows[-1] += [""] * (3 - len(rows[-1]))
        pt = Table(rows, colWidths=[col_w] * 3)
        pt.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
        story.append(pt)

    doc.build(story)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="ceklis-{tx.plate.replace(" ", "")}-{tx.id}.pdf"'})
